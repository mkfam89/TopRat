#!/usr/bin/env python3
"""Build the synthetic lint_resume.py fixtures in tests/fixtures/docx/.

    python tests/build_docx_fixtures.py

Each file violates exactly ONE rule (R01..R12); clean_pass.docx violates none.
R02/R03 live in the PDF side, so their .docx content is clean — the tests
monkeypatch lint_resume.pdf_pages to simulate page counts (offline, no
LibreOffice/pdftotext needed).
"""
import os
import sys

import docx
from docx.shared import Pt

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures', 'docx')
os.makedirs(OUT, exist_ok=True)

DATE = '01/2021 - Present'
OLD_DATE = '03/2016 - 12/2018'


def header(d, text):
    d.add_paragraph(text)          # uppercase short paragraph => section header


def role(d, company, title, date=DATE):
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = company          # company row (no date)
    t.rows[0].cells[1].text = 'Houston, TX'
    t2 = d.add_table(rows=1, cols=2)
    t2.rows[0].cells[0].text = title           # role row (right cell has a date)
    t2.rows[0].cells[1].text = date


def bullets(d, n, seed='Automated deployment pipelines cutting release time'):
    for i in range(n):
        d.add_paragraph(f'• {seed} across environment {i + 1} of the platform.')


def skeleton(summary_opener='Site reliability engineer',
             skills_line='Cloud - Azure, Kubernetes, Docker, Terraform tooling',
             n_bullets=2):
    d = docx.Document()
    d.add_paragraph('Jane Doe')
    header(d, 'SUMMARY')
    d.add_paragraph(f'{summary_opener} with eight years of experience running '
                    'production platforms and cutting operational toil at scale.')
    header(d, 'TECHNICAL SKILLS')
    d.add_paragraph(skills_line)
    d.add_paragraph('Monitoring - Prometheus, Grafana, PagerDuty alerting stack')
    header(d, 'EXPERIENCE')
    role(d, 'Acme Corp', 'Site Reliability Engineer')
    bullets(d, n_bullets)
    header(d, 'EDUCATION')
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = 'BS Computer Science'
    t.rows[0].cells[1].text = 'University of Houston'
    return d


def save(d, name):
    path = os.path.join(OUT, name)
    d.save(path)
    print('wrote', name)


def main():
    # clean — must produce zero findings (docx side)
    save(skeleton(), 'clean_pass.docx')

    # R01 font floor: one 9pt run
    d = skeleton()
    p = d.add_paragraph()
    p.add_run('Fine print line that dips below the ten point floor.').font.size = Pt(9)
    save(d, 'r01_font_min.docx')

    # R02 / R03 are PDF-side: clean docx content, page counts faked in the test
    save(skeleton(), 'r02_page_count.docx')
    save(skeleton(), 'r03_page2_lines.docx')

    # R04 one-bullet role
    save(skeleton(n_bullets=1), 'r04_bullets_per_role.docx')

    # R05 zero-bullet role
    save(skeleton(n_bullets=0), 'r05_orphan_title.docx')

    # R06 banned header
    d = skeleton()
    header(d, 'CORE COMPETENCIES')
    d.add_paragraph('Leadership, communication, collaboration and planning strengths.')
    save(d, 'r06_banned_header.docx')

    # R07 required section missing (no TECHNICAL SKILLS)
    d = docx.Document()
    d.add_paragraph('Jane Doe')
    header(d, 'SUMMARY')
    d.add_paragraph('Site reliability engineer with eight years of experience '
                    'running production platforms and cutting operational toil.')
    header(d, 'EXPERIENCE')
    role(d, 'Acme Corp', 'Site Reliability Engineer')
    bullets(d, 2)
    save(d, 'r07_required_section.docx')

    # R08 notes leakage
    d = skeleton()
    d.add_paragraph('Note: verify the salary range before applying to this one.')
    save(d, 'r08_notes_leakage.docx')

    # R09 summary adopts the target job title (filename carries the role slug)
    save(skeleton(summary_opener='Data engineer'),
         'Jane_Doe_Resume_Acme_DataEngineer.docx')

    # R10 missing pdf twin — clean content, tested with check_pdf=True
    save(skeleton(), 'r10_missing_pdf.docx')

    # R11 double colon in a skills line
    save(skeleton(skills_line='Cloud :: Azure, Kubernetes, Docker, Terraform tooling'),
         'r11_skills_format.docx')

    # R12 trailing role appears alone (1 of the group of 3)
    d = skeleton()
    role(d, 'BigCo', 'Senior Implementation Consultant', OLD_DATE)
    bullets(d, 2, seed='Led onboarding programs raising integration throughput')
    save(d, 'r12_trailing_group.docx')


if __name__ == '__main__':
    sys.exit(main())
